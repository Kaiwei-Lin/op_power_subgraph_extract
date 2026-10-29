from contextlib import redirect_stdout
import csv
from io import StringIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from onnx_graph import merge_power as merge
from onnx_graph.compiled_stats import number_models
from onnx_graph.omc import _schema
from onnx_graph.tests import test_compiled_stats as fixtures


class BatchPowerTests(unittest.TestCase):
    save_model = fixtures.CompiledStatisticsTests.save_model

    @classmethod
    def setUpClass(cls):
        try:
            cls.Model, _, _, cls.schema_path = _schema()
        except ValueError as error:
            raise unittest.SkipTest(str(error)) from error

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="batch-power-test-", dir=Path.cwd())
        self.root = Path(self.temp.name).resolve()
        self.assertEqual(self.root.parent, Path.cwd().resolve())
        self.addCleanup(self.temp.cleanup)
        self.models = self.root / "batch"
        self.models.mkdir()
        self.powers = self.root / "power"
        self.powers.mkdir()
        self.output = self.root / "result"

    def power_csv(self, number, values=None, db="Work0.db", encoding="utf-8-sig"):
        values = values or {"CPU": "0", "NPU": "3.25 mW"}
        path = self.powers / f"{number}.csv"
        with path.open("w", encoding=encoding, newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["metadata"])
            writer.writerow(["DB", "Inst", "Channel", "Current", "Voltage", "Power"])
            writer.writerows([[db, "inst", c, "unused", "unused", v] for c, v in values.items()])
        return path

    def run_merge(self, **kwargs):
        return merge.merge_compiled_power(self.models, self.powers, self.output,
                                           ge_proto_path=self.schema_path, **kwargs)

    def set_csv_runtime(self, path, value):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        rows[0] = ["metadata", "", "", "", "", "", "not the duration", value]
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(rows)

    def test_numbering_uses_numeric_slash_and_case_insensitive_names(self):
        names = ["slash_2/sg_a.omc", "slash_10/sg_a.omc", "slash0/sg_Gemm.omc",
                 "slash0/sg_encoder.omc", "slash0/sg_decoder.omc", "slash_1/sg_a.omc"]
        for name in names:
            self.save_model(name)
        rows = number_models(self.models)
        self.assertEqual([r["OMC相对路径"] for r in rows], [names[i] for i in (4, 3, 2, 5, 0, 1)])
        self.assertEqual([r["编号"] for r in rows], list(range(1, 7)))

    def test_generic_batch_does_not_use_fixed_legacy_map_and_keeps_power_units(self):
        self.save_model("slash_0/sg_decoder.omc")
        self.save_model("slash_1/sg_encoder.omc")
        first = self.power_csv(1)
        second = self.power_csv(2, {"NPU": "2.5 mW", "CPU": "7"})
        before = first.read_bytes(), second.read_bytes()
        with patch.object(merge, "MODEL_FILES", {}):
            result = self.run_merge()
        self.assertEqual(result["omc_counts"], {"ok": 2})
        self.assertEqual(result["power_counts"], {"ok": 2})
        self.assertEqual(result["records"][0]["CPU"], "0")
        self.assertEqual(result["records"][0]["NPU"], "3.25 mW")
        self.assertEqual(result["records"][1]["NPU"], "2.5 mW")
        for row in result["records"]:
            self.assertEqual((row["OMC_节点数"], row["OMC_Add数量"], row["OMC_FLOPs"], row["OMC_数据转移字节数"]), (2, 2, 8, 96))
        with (self.output / "summary.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([r["编号"] for r in rows], ["1", "2"])
        self.assertEqual((first.read_bytes(), second.read_bytes()), before)

    def test_saved_numbering_validates_hashes_inventory_and_ids(self):
        model = self.save_model("slash_0/sg_a.omc")
        result = merge.merge_compiled_power(self.models, None, self.output, ge_proto_path=self.schema_path)
        self.assertEqual(result["power_counts"], {"missing": 1})
        numbering = self.output / "numbering.csv"
        self.assertEqual(number_models(self.models, numbering)[0]["编号"], 1)
        self.power_csv(1)
        merged = merge.merge_compiled_power(self.models, self.powers, self.root / "with_power",
                                            ge_proto_path=self.schema_path, numbering_csv=numbering)
        self.assertEqual(merged["power_counts"], {"ok": 1})
        raw = model.read_bytes()
        model.write_bytes(raw + b"changed")
        with self.assertRaisesRegex(ValueError, "内容或层级"):
            number_models(self.models, numbering)
        model.write_bytes(raw)
        self.save_model("slash_0/sg_new.omc")
        with self.assertRaisesRegex(ValueError, "完整覆盖"):
            number_models(self.models, numbering)

    def test_missing_failed_ambiguous_and_incomplete_power_are_visible(self):
        for i in range(5):
            self.save_model(f"slash_0/sg_{i}.omc")
        self.power_csv(1, {"CPU": "0", "NPU": "1"})
        self.power_csv(2, {"CPU": "2"})
        (self.powers / "3.csv").write_text("not a power table", encoding="utf-8")
        self.power_csv(4)
        (self.powers / "4.xlsx").write_bytes(b"both formats")
        self.power_csv(99)
        result = self.run_merge()
        self.assertEqual(result["power_counts"], {"ok": 1, "incomplete": 1, "failed": 1, "ambiguous": 1, "missing": 1})
        self.assertEqual(result["records"][-1]["OMC状态"], "ok")
        self.assertNotIn("CPU", result["records"][-1])
        self.assertEqual(len(result["unused_power_files"]), 1)

    def test_csv_prefers_work0_and_rejects_duplicate_channels(self):
        path = self.power_csv(1, {"CPU": "2"}, db="Work1.db", encoding="gb18030")
        with path.open("a", encoding="gb18030", newline="") as stream:
            csv.writer(stream).writerows([["DB", "Inst", "Channel", "Current", "Voltage", "Power"],
                                         ["Work0.db", "inst", "CPU", "", "", "0 毫瓦"]])
        self.assertEqual(merge.read_power_csv(path), {"CPU": "0 毫瓦"})
        path.write_text("Channel,Power\nCPU,1\nCPU,2\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "重复"):
            merge.read_power_csv(path)

    def test_csv_uses_first_twelve_channel_rows_like_the_old_workbook_reader(self):
        path = self.power_csv(1, {f"CH_{i}": str(i) for i in range(13)})
        powers = merge.read_power_csv(path)
        self.assertEqual(len(powers), 12)
        self.assertNotIn("CH_12", powers)

    def test_h1_runtime_and_transfer_throughput_keep_zero_and_missing_distinct(self):
        for i in range(3):
            self.save_model(f"slash_0/sg_{i}.omc")
            path = self.power_csv(i + 1)
            if i < 2:
                self.set_csv_runtime(path, "0.5" if i == 0 else "0")
        result = self.run_merge()
        self.assertEqual(result["runtime_counts"], {"ok": 2, "missing": 1})
        first, zero, missing = result["records"]
        self.assertEqual(first["运行时长"], 0.5)
        self.assertEqual(first["数据移动吞吐"], 192)
        self.assertEqual((first["运行时长单位"], first["数据移动吞吐单位"]), ("s", "B/s"))
        self.assertEqual(zero["运行时长"], 0)
        self.assertIsNone(zero["数据移动吞吐"])
        self.assertIsNone(missing["运行时长"])
        self.assertIsNone(missing["数据移动吞吐"])
        with (self.output / "summary.csv").open(encoding="utf-8-sig", newline="") as stream:
            saved = list(csv.DictReader(stream))
        self.assertEqual(saved[0]["数据移动吞吐"], "192.0")
        self.assertEqual(saved[1]["数据移动吞吐"], "")
        self.assertEqual(saved[2]["运行时长"], "")

    def test_bad_runtime_does_not_discard_power_and_reads_physical_first_csv_row(self):
        self.save_model("slash_0/sg_a.omc")
        path = self.power_csv(1)
        for value in ("-1", "NaN", "inf", "invalid", "1e999", "1e-999"):
            with self.subTest(value=value):
                self.set_csv_runtime(path, value)
                with self.assertRaises(ValueError):
                    merge.read_runtime_seconds(path)
        result = self.run_merge()
        row = result["records"][0]
        self.assertEqual(row["功耗状态"], "ok")
        self.assertEqual(row["NPU"], "3.25 mW")
        self.assertEqual(row["时长状态"], "failed")
        self.assertIsNone(row["数据移动吞吐"])
        self.set_csv_runtime(path, "1e-3")
        self.assertEqual(merge.read_runtime_seconds(path), .001)
        content = path.read_text(encoding="utf-8-sig")
        path.write_text("\n" + content, encoding="utf-8-sig")
        self.assertIsNone(merge.read_runtime_seconds(path))

    def test_incomplete_transfer_is_not_used_for_throughput(self):
        self.save_model(unknown=True)
        path = self.power_csv(1)
        self.set_csv_runtime(path, "2")
        row = self.run_merge()["records"][0]
        self.assertEqual(row["运行时长"], 2)
        self.assertIsNone(row["OMC_数据转移字节数"])
        self.assertGreater(row["OMC_已知数据转移字节数"], 0)
        self.assertIsNone(row["数据移动吞吐"])

    def test_xlsx_runtime_uses_h1_from_the_selected_average_sheet(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("XLSX 时长测试需要 openpyxl")
        self.save_model("slash_0/sg_a.omc")
        path = self.powers / "1.xlsx"
        book = Workbook()
        book.active.title = "other"
        book.active["H1"] = 999
        sheet = book.create_sheet(merge.SHEET_NAME)
        sheet["H1"] = 4
        sheet.append(["DB", "Inst", "Channel", "Current", "Voltage", "Power"])
        sheet.append(["Work0.db", "inst", "CPU", 1, 1, 2.5])
        book.save(path)
        book.close()
        row = self.run_merge()["records"][0]
        self.assertEqual((row["运行时长"], row["数据移动吞吐"], row["CPU"]), (4, 24, 2.5))
        book = load_workbook(path)
        book[merge.SHEET_NAME]["H1"] = "#VALUE!"
        book.save(path)
        book.close()
        with self.assertRaises(ValueError):
            merge.read_runtime_seconds(path)

    def test_cli_and_existing_output_are_safe(self):
        self.save_model("slash_0/sg_a.omc")
        self.power_csv(1)
        args = ["merge_power", "--batch-dir", str(self.models), "--power-dir", str(self.powers),
                "--output-dir", str(self.output), "--ge-proto", self.schema_path]
        with patch("sys.argv", args), redirect_stdout(StringIO()) as printed:
            self.assertEqual(merge.main(), 0)
        self.assertIn("numbering.csv", printed.getvalue())
        before = (self.output / "summary.csv").read_bytes()
        with self.assertRaises(FileExistsError):
            self.run_merge()
        self.assertEqual((self.output / "summary.csv").read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
