import csv
from copy import deepcopy
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

try:
    from openpyxl import Workbook, load_workbook
except ImportError as error:
    raise unittest.SkipTest("Excel 功耗测试需要 openpyxl；CSV/OMC 统计不依赖它") from error

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import merge_power


CHANNELS = ["VDD_CPU_L", "VDD_CPU_M", "VDD_CPU_B", "VDD_CPUM_MEM"] + [f"CH_{i}" for i in range(8)]
STAT_FIELDS = ["子图名字", "节点数", "Conv数量", "MACs", "FLOPs", "MFLOPs", "计算量完整",
               "未知计算量节点数", "已知MACs", "已知FLOPs", "exec_time", "exec_time_unit",
               "算术强度", "算术强度单位", "已知exec_time", "时间匹配节点数", "时间未匹配节点数"]


def write_measurement(path, number):
    book = Workbook()
    book.active.title = "Average Values Work0.db info"
    book.active.append(["ignore", 999999])
    sheet = book.create_sheet(merge_power.SHEET_NAME)
    sheet.append(["metadata", None, "IGNORE_METADATA", None, None, 999999])
    sheet.append(["DB", "Inst", "Channel", "Current", "Voltage", "Power"])
    for index in range(12):
        channel_index = (index + number) % 12
        value = number * 100 + channel_index + 0.125
        if channel_index == 0:
            value = 0
        if channel_index == 11:
            value = f"{value} mW"
        sheet.append(["Work0.db", "inst", CHANNELS[channel_index], 999, 888, value])
    sheet.append([])
    sheet.append(["extra", None, "IGNORE_EXTRA", None, None, 999999])
    book.save(path)
    book.close()


class MergePowerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="power-merge-test-", dir=Path.cwd())
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.assertEqual(self.root.parent, Path.cwd().resolve())
        self.power_dir = self.root / "功耗 表"
        self.power_dir.mkdir()
        self.summary = self.root / "representatives_summary.csv"
        self.output = self.root / "result.csv"
        self.fields = ["斜杠数", "子图名字", "ONNX文件", "节点数", "exec_time"]
        self.rows = []
        for number in range(56, 0, -1):
            model = merge_power.MODEL_FILES[number]
            folder = model.split("/")[0]
            stem = merge_power.model_key(model)
            # 验证原始 ONNX、修复 ONNX 和编译 OMC，以及 Windows/Linux 路径。
            filename = stem + ("_fixed.onnx" if number % 3 == 0 else ".onnx")
            path = f"D:\\公司 数据\\{folder}\\{filename}"
            if number % 3 == 1:
                path = "/company/" + model
            prefix = "/".join(["层级"] * int(folder[-1]) + [f"子图,{number}"])
            self.rows.append({"斜杠数": folder[-1], "子图名字": prefix,
                              "ONNX文件": path, "节点数": "003", "exec_time": "0.00100"})
        self.save_summary()

    def save_summary(self):
        with self.summary.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=self.fields)
            writer.writeheader()
            writer.writerows(self.rows)

    def write_statistics(self):
        stats_dir = self.root / "handfree_ainr_stats"
        stats_dir.mkdir(exist_ok=True)
        self.expected_stats = {}
        nodes, levels = [], []
        for depth in range(4):
            records = []
            for row in self.rows:
                if row["斜杠数"] != str(depth):
                    continue
                number = int(row["子图名字"].rsplit(",", 1)[1])
                flops = 0 if number == 2 else number * 1000
                record = {
                    "子图名字": "/" + row["子图名字"].replace("/", "//") + "/",
                    "节点数": "3", "Conv数量": "3", "MACs": str(flops // 2), "FLOPs": str(flops),
                    "MFLOPs": str(flops / 1e6), "计算量完整": "是", "未知计算量节点数": "0",
                    "已知MACs": str(flops // 2), "已知FLOPs": str(flops), "exec_time": "0.001",
                    "exec_time_unit": "us", "算术强度": str(flops / 0.001), "算术强度单位": "FLOPs/us",
                    "已知exec_time": "0.001", "时间匹配节点数": "3", "时间未匹配节点数": "0",
                }
                if number == 19:
                    record.update({field: "" for field in ("MACs", "FLOPs", "MFLOPs", "算术强度")})
                    record.update({"计算量完整": "否", "未知计算量节点数": "1"})
                self.expected_stats[row["子图名字"]] = dict(record)
                records.append(record)
            # 原统计包含未入选代表清单的子图；不应添加到最终表。
            extra = dict(records[0], 子图名字="/".join(["未入选"] * (depth + 1)))
            with (stats_dir / f"stats_{depth}_slash.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=STAT_FIELDS)
                writer.writeheader()
                writer.writerows([extra, *reversed(records)])
            groups = []
            for record in [extra, *records]:
                ids = list(range(len(nodes), len(nodes) + 3))
                for node_id in ids:
                    nodes.append({"node_id": node_id,
                                  "inputs": [{"name": "x", "shape": [2, 3], "dtype": 1},
                                             {"name": "weight", "shape": [3, 4], "dtype": 10}],
                                  "outputs": [{"name": "y", "shape": [2, 4], "dtype": 1}]})
                groups.append({"prefix": record["子图名字"], "node_count": 3, "node_ids": ids})
            levels.append({"slash_count": depth, "subgraphs": groups})
        total = {
            "node_count": len(nodes), "operator_counts": {"Conv": len(nodes)},
            "macs": 98765, "flops": 987654, "complete": True, "unknown_node_ids": [],
            "known_macs": 98765, "known_flops": 987654, "exec_time": 10,
            "arithmetic_intensity": 98765.4, "known_exec_time": 10,
            "exec_time_matched_nodes": len(nodes), "exec_time_missing_node_ids": [],
        }
        (stats_dir / "stats.json").write_text(
            json.dumps({"source": "D:/models/handfree_ainr.onnx", "total": total,
                        "nodes": nodes, "levels": levels}), encoding="utf-8")
        return stats_dir

    def test_cli_merges_all_57_files_without_reordering_or_changing_originals(self):
        for number in range(57):
            write_measurement(self.power_dir / f"{number}.xlsx", number)
        before = {p: p.read_bytes() for p in [self.summary, *self.power_dir.glob("*.xlsx")]}
        result = subprocess.run(
            [sys.executable, merge_power.__file__, "--summary", str(self.summary),
             "--power-dir", str(self.power_dir), "-o", str(self.output)],
            capture_output=True, text=True, encoding="utf-8",
            env={**os.environ, "PYTHONIOENCODING": "utf-8"})
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("56 个子图 + 1 个整模型", result.stdout)
        self.assertTrue(self.output.read_bytes().startswith(b"\xef\xbb\xbf"))
        with self.output.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            merged = list(reader)
            self.assertEqual(reader.fieldnames, self.fields + CHANNELS)
        self.assertEqual(len(merged), 57)
        for original, row, number in zip(self.rows, merged, range(56, 0, -1)):
            self.assertEqual({key: row[key] for key in self.fields}, original)
            for channel_index, channel in enumerate(CHANNELS):
                expected = str(number * 100 + channel_index + 0.125)
                if channel_index == 0:
                    expected = "0"
                if channel_index == 11:
                    expected += " mW"
                self.assertEqual(row[channel], expected)
        self.assertEqual(merged[-1]["子图名字"], "整模型")
        self.assertTrue(all(merged[-1][field] == "" for field in self.fields if field != "子图名字"))
        self.assertEqual(merged[-1]["VDD_CPU_L"], "0")
        self.assertEqual(merged[-1]["VDD_CPU_M"], "1.125")
        for path, content in before.items():
            self.assertEqual(path.read_bytes(), content)

    def test_fixed_suffix_and_hash_are_preserved(self):
        self.assertEqual(merge_power.model_key(merge_power.MODEL_FILES[5]), "sg_dual_model_f3db236e3a66")
        self.assertEqual(merge_power.model_key(merge_power.MODEL_FILES[27]),
                         "sg_iccrn_model_oSpatialNet_freq_cmps_1_3b2388043deb")
        self.assertNotEqual(merge_power.model_key("sg_Concat_1_58175fb15a27.onnx"),
                            merge_power.model_key("sg_Concat_1_000000000000.onnx"))

    def test_cli_merges_statistics_and_full_model_totals(self):
        stats_dir = self.write_statistics()
        for number in range(57):
            write_measurement(self.power_dir / f"{number}.xlsx", number)
        inputs = [self.summary, *stats_dir.iterdir(), *self.power_dir.iterdir()]
        before = {path: path.read_bytes() for path in inputs}
        result = subprocess.run(
            [sys.executable, merge_power.__file__, "--summary", str(self.summary),
             "--power-dir", str(self.power_dir), "--stats-dir", str(stats_dir)],
            capture_output=True, text=True, encoding="utf-8",
            env={**os.environ, "PYTHONIOENCODING": "utf-8"})
        self.assertEqual(result.returncode, 0, result.stderr)
        output = self.root / "representatives_summary_with_power_stats.csv"
        added_fields = [field for field in STAT_FIELDS if field not in self.fields]
        with output.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            merged = list(reader)
            self.assertEqual(reader.fieldnames, self.fields + added_fields + merge_power.TRANSFER_FIELDS + CHANNELS)
        self.assertEqual(len(merged), 57)
        for original, row in zip(self.rows, merged):
            self.assertEqual({key: row[key] for key in self.fields}, original)
            for field in added_fields:
                self.assertEqual(row[field], self.expected_stats[original["子图名字"]][field])
            self.assertEqual(row["VDD_CPU_L"], "0")
            # 三个算子，每个读取 FP32[2,3] + FP16[3,4]，写入 FP32[2,4]。
            self.assertEqual(row["输入及权重字节数"], "144")
            self.assertEqual(row["输出字节数"], "96")
            self.assertEqual(row["数据转移字节数"], "240")
            self.assertEqual(row["数据转移量完整"], "是")
        self.assertEqual(merged[-1]["子图名字"], "整模型")
        self.assertEqual(merged[-1]["ONNX文件"], "D:/models/handfree_ainr.onnx")
        self.assertEqual(merged[-1]["节点数"], "180")
        self.assertEqual(merged[-1]["Conv数量"], "180")
        self.assertEqual(merged[-1]["FLOPs"], "987654")
        self.assertEqual(merged[-1]["算术强度"], "98765.4")
        self.assertEqual(merged[-1]["算术强度单位"], "FLOPs/us")
        self.assertEqual(merged[-1]["VDD_CPU_M"], "1.125")
        self.assertEqual(merged[-1]["数据转移字节数"], "14400")
        for path, content in before.items():
            self.assertEqual(path.read_bytes(), content)

    def test_missing_duplicate_conflicting_or_invalid_statistics_fail_before_output(self):
        for case in ("missing", "duplicate", "conflict", "wrong_depth", "missing_flops", "bad_json"):
            with self.subTest(case=case):
                stats_dir = self.write_statistics()
                path = stats_dir / "stats_3_slash.csv"
                columns, records = merge_power.read_csv(path, ())
                if case == "missing":
                    records.pop()
                elif case == "duplicate":
                    records.append(dict(records[-1]))
                elif case == "conflict":
                    records[-1]["exec_time"] = "999"
                elif case == "wrong_depth":
                    records[-1]["子图名字"] = "wrong/depth"
                elif case == "missing_flops":
                    columns.remove("FLOPs")
                else:
                    (stats_dir / "stats.json").write_text("{}", encoding="utf-8")
                with path.open("w", encoding="utf-8-sig", newline="") as stream:
                    writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
                    writer.writeheader()
                    writer.writerows(records)
                with self.assertRaises(ValueError):
                    merge_power.merge_power(self.summary, self.power_dir, self.output, stats_dir=stats_dir)
                self.assertFalse(self.output.exists())

    def test_missing_intensity_is_reported_without_inventing_values(self):
        stats_dir = self.write_statistics()
        removed = {"exec_time", "exec_time_unit", "算术强度", "算术强度单位", "已知exec_time",
                   "时间匹配节点数", "时间未匹配节点数"}
        for path in stats_dir.glob("*.csv"):
            columns, records = merge_power.read_csv(path, ())
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=[c for c in columns if c not in removed],
                                        extrasaction="ignore")
                writer.writeheader()
                writer.writerows(records)
        fields, rows = merge_power.read_csv(self.summary, ())
        errors = io.StringIO()
        with patch("sys.stderr", errors):
            columns, total = merge_power.merge_statistics(fields, rows, stats_dir)
        self.assertNotIn("算术强度", columns)
        self.assertEqual(total["数据转移字节数"], 14400)
        self.assertIn("没有算术强度列", errors.getvalue())
        self.assertEqual(rows[0]["exec_time"], "0.00100")

    def test_missing_json_is_reported_as_required_tensor_metadata(self):
        stats_dir = self.write_statistics()
        (stats_dir / "stats.json").unlink()
        with self.assertRaisesRegex(FileNotFoundError, "完整 stats.json"):
            merge_power.merge_power(self.summary, self.power_dir, self.output, stats_dir=stats_dir)
        self.assertFalse(self.output.exists())

    def test_transfer_counts_weight_reads_and_internal_edges_without_summing_levels(self):
        nodes = [{"node_id": number, "flops": None,
                  "inputs": [{"name": "x" if number == 0 else "mid", "shape": [4], "dtype": 1},
                             {"name": "shared_weight", "shape": [2, 2], "dtype": 10}],
                  "outputs": [{"name": "mid" if number == 0 else "y", "shape": [4], "dtype": 1}]}
                 for number in range(2)]
        report = {"nodes": nodes, "total": {"node_count": 2}, "levels": [
            {"slash_count": 0, "subgraphs": [{"prefix": "root", "node_ids": [0, 1], "node_count": 2}]},
            {"slash_count": 1, "subgraphs": [
                {"prefix": "root/a", "node_ids": [0], "node_count": 1},
                {"prefix": "root/b", "node_ids": [1], "node_count": 1}]}]}
        groups, total = merge_power.transfer_statistics(report)
        self.assertEqual(groups[(1, "root/a")]["数据转移字节数"], 40)
        self.assertEqual(groups[(1, "root/b")]["数据转移字节数"], 40)
        self.assertEqual(groups[(0, "root")]["数据转移字节数"], 80)
        self.assertEqual(total["输入及权重字节数"], 48)
        self.assertEqual(total["输出字节数"], 32)
        self.assertEqual(total["数据转移字节数"], 80)
        self.assertEqual(total["数据转移量完整"], "是")

    def test_unknown_transfer_keeps_known_part_and_zero_sized_tensors(self):
        report = {"nodes": [
            {"node_id": 0, "inputs": [{"shape": ["batch", 4], "dtype": 1}],
             "outputs": [{"shape": [8], "dtype": 10}]},
            {"node_id": 1, "inputs": [], "outputs": [{"shape": [0, 100], "dtype": 1}]},
        ], "total": {"node_count": 2}, "levels": []}
        _, total = merge_power.transfer_statistics(report)
        self.assertIsNone(total["输入及权重字节数"])
        self.assertIsNone(total["数据转移字节数"])
        self.assertEqual(total["输出字节数"], 16)
        self.assertEqual(total["已知数据转移字节数"], 16)
        self.assertEqual(total["未知数据转移节点数"], 1)
        self.assertEqual(total["数据转移量完整"], "否")

    def test_tensor_byte_width_scalar_packed_and_unknown_types(self):
        cases = [(1, [3], 12), (10, [3], 6), (16, [3], 6), (7, [3], 24),
                 (9, [3], 3), (1, [], 4), (22, [3], 2), (14, [3], 24),
                 ("DT_FLOAT16", [3], 6), (1, [0, 3], 0),
                 (1, [None, 3], None), (1, [-1], None), (1, None, None),
                 (1, [True], None), (8, [3], None), (999, [3], None)]
        for dtype, shape, expected in cases:
            with self.subTest(dtype=dtype, shape=shape):
                self.assertEqual(merge_power.tensor_bytes({"shape": shape, "dtype": dtype}), expected)

    def test_transfer_rejects_incomplete_or_ambiguous_membership_metadata(self):
        stats_dir = self.write_statistics()
        original = json.loads((stats_dir / "stats.json").read_text(encoding="utf-8"))
        for case in ("no_nodes", "duplicate_node", "missing_member", "duplicate_member", "no_outputs"):
            with self.subTest(case=case):
                report = deepcopy(original)
                if case == "no_nodes":
                    report.pop("nodes")
                elif case == "duplicate_node":
                    report["nodes"][1]["node_id"] = 0
                elif case == "missing_member":
                    report["levels"][0]["subgraphs"][0]["node_ids"][0] = 99999
                elif case == "duplicate_member":
                    report["levels"][0]["subgraphs"][0]["node_ids"][0] = 1
                else:
                    report["nodes"][0].pop("outputs")
                with self.assertRaises(ValueError):
                    merge_power.transfer_statistics(report)

    def test_unknown_full_model_flops_and_intensity_remain_missing(self):
        stats_dir = self.write_statistics()
        path = stats_dir / "stats.json"
        report = json.loads(path.read_text(encoding="utf-8"))
        report["total"].update(macs=None, flops=None, complete=False, unknown_node_ids=[1],
                               exec_time=None, arithmetic_intensity=None, exec_time_missing_node_ids=[1])
        path.write_text(json.dumps(report), encoding="utf-8")
        total = merge_power.full_model_statistics(stats_dir, STAT_FIELDS)
        for field in ("MACs", "FLOPs", "MFLOPs", "exec_time", "算术强度"):
            self.assertIsNone(total[field])
        self.assertEqual(total["计算量完整"], "否")
        self.assertEqual(total["未知计算量节点数"], 1)
        self.assertEqual(total["已知FLOPs"], 987654)

    def test_bad_measurements_are_rejected(self):
        path = self.power_dir / "1.xlsx"
        for case in ("duplicate", "blank", "formula", "excel_error", "header", "sheet"):
            with self.subTest(case=case):
                write_measurement(path, 1)
                book = load_workbook(path)
                sheet = book[merge_power.SHEET_NAME]
                if case == "duplicate":
                    sheet["C4"] = sheet["C3"].value
                elif case == "blank":
                    sheet["F3"] = None
                elif case == "formula":
                    sheet["F3"] = "=1+1"
                elif case == "excel_error":
                    sheet["F3"] = "#VALUE!"
                elif case == "header":
                    sheet["F2"] = "Wrong"
                else:
                    sheet.title = "Wrong sheet"
                book.save(path)
                book.close()
                with self.assertRaises(ValueError):
                    merge_power.read_power(path)

    def test_work1_is_used_only_when_work0_is_missing(self):
        path = self.power_dir / "1.xlsx"
        write_measurement(path, 1)
        book = load_workbook(path)
        fallback = book.copy_worksheet(book[merge_power.SHEET_NAME])
        fallback.title = merge_power.FALLBACK_SHEET_NAME
        fallback["F3"] = 9999
        book.save(path)
        self.assertEqual(merge_power.read_power(path)["VDD_CPU_M"], 101.125)

        book[merge_power.SHEET_NAME]["F2"] = "Wrong"
        book.save(path)
        with self.assertRaises(ValueError):
            merge_power.read_power(path)

        del book[merge_power.SHEET_NAME]
        book.save(path)
        book.close()
        self.assertEqual(merge_power.read_power(path)["VDD_CPU_M"], 9999)

    def test_unmatched_duplicate_and_missing_rows_are_rejected_before_output(self):
        original = list(self.rows)
        for case in ("unknown", "duplicate", "missing", "ragged"):
            with self.subTest(case=case):
                self.rows = [dict(row) for row in original]
                if case == "unknown":
                    self.rows[0]["ONNX文件"] = "unrelated_model.onnx"
                elif case == "duplicate":
                    self.rows.append(dict(self.rows[0]))
                elif case == "missing":
                    self.rows.pop()
                self.save_summary()
                if case == "ragged":
                    with self.summary.open("a", encoding="utf-8", newline="") as stream:
                        stream.write("too,few,fields\n")
                with self.assertRaises(ValueError):
                    merge_power.merge_power(self.summary, self.power_dir, self.output)
                self.assertFalse(self.output.exists())

    def test_missing_xlsx_does_not_create_output(self):
        with self.assertRaisesRegex(FileNotFoundError, "0.xlsx"):
            merge_power.merge_power(self.summary, self.power_dir, self.output)
        self.assertFalse(self.output.exists())

    def test_existing_output_and_input_are_not_overwritten(self):
        self.output.write_text("keep", encoding="utf-8")
        original = self.summary.read_bytes()
        for output in (self.output, self.summary):
            with self.assertRaises(FileExistsError):
                merge_power.merge_power(self.summary, self.power_dir, output)
        self.assertEqual(self.output.read_text(encoding="utf-8"), "keep")
        self.assertEqual(self.summary.read_bytes(), original)

    def test_channel_mismatch_and_existing_columns_are_rejected(self):
        for number in range(57):
            (self.power_dir / f"{number}.xlsx").touch()
        for case in ("missing_channel", "existing_column"):
            with self.subTest(case=case):
                if case == "existing_column":
                    self.fields.append("VDD_CPU_L")
                    self.save_summary()
                def measured(path):
                    if case == "missing_channel" and path.stem == "56":
                        return {"VDD_CPU_M": 1}
                    return {"VDD_CPU_L": 0, "VDD_CPU_M": 1}
                with patch.object(merge_power, "read_power", side_effect=measured):
                    with self.assertRaises(ValueError):
                        merge_power.merge_power(self.summary, self.power_dir, self.output)
                self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
