#!/usr/bin/env python3
"""通用 OMC 编号、模型统计与同号功耗 CSV 合并。

通用模式（不使用固定 MODEL_FILES 映射）：
    python merge_power.py --batch-dir D:/models/batch --power-dir D:/power --output-dir D:/results/model_power
先统计并保存编号（功耗尚未提供）：
    python merge_power.py --batch-dir D:/models/batch --output-dir D:/results/model_stats
后续增加 --numbering D:/results/model_stats/numbering.csv 固定编号并验证 OMC SHA256。
通用模式需要同目录的 compiled_stats.py、stats.py、omc.py 等解析模块。
每个 OMC 按整个编译后运算图统计，输出算子数量、理论 FLOPs 和逐算子数据转移量。
CSV 功耗读取 Channel/Power，取表头后前 12 条非空数据行，保留原值及单位。
运行时长读取功耗文件 H1，单位 s；数据移动吞吐=完整转移字节数/时长，单位 B/s。
时长缺失、无效、为零或数据量不完整时，吞吐量保持空白。
兼容同号 XLSX；仅读取 Excel 时需要 python -m pip install openpyxl。

旧 --summary 模式：
    python merge_power.py --summary D:/test/representatives_summary.csv --power-dir D:/test/power --stats-dir D:/test/model_stats
旧模式兼容既有的固定功耗编号、0.xlsx 整模型及 --stats-dir 统计合并。
Excel 优先 Average Values Work0.db，其次 Average Values Work1.db。
读取第 2 行表头及第 3～14 行数据；不换算功耗单位，同名列冲突时报错。
所有模式均保留源模型及功耗文件，输出位置必须尚不存在。
数据转移量 = 每个运算节点输入（包含权重/偏置）及输出字节数之和。
中间张量的写入和读取分别计入；动态形状或未知类型保留已知部分并标注不完整。
"""

import argparse
import csv
from decimal import Decimal, InvalidOperation
import json
import math
import re
import sys
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile

try:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError:
    load_workbook = None
    InvalidFileException = ValueError


SHEET_NAME = "Average Values Work0.db"
FALLBACK_SHEET_NAME = "Average Values Work1.db"
TRANSFER_FIELDS = ["输入及权重字节数", "输出字节数", "数据转移字节数", "数据转移量完整",
                   "未知数据转移节点数", "已知数据转移字节数"]
# ONNX TensorProto 的固定宽度类型；STRING 和未支持的类型保持未知。
DTYPE_BITS = {1: 32, 2: 8, 3: 8, 4: 16, 5: 16, 6: 32, 7: 64, 9: 8, 10: 16,
              11: 64, 12: 32, 13: 64, 14: 64, 15: 128, 16: 16,
              17: 8, 18: 8, 19: 8, 20: 8, 21: 4, 22: 4, 23: 4, 24: 8}
DTYPE_BITS.update({
    "DT_FLOAT": 32, "DT_FLOAT16": 16, "DT_DOUBLE": 64, "DT_BF16": 16,
    "DT_INT8": 8, "DT_UINT8": 8, "DT_INT16": 16, "DT_UINT16": 16,
    "DT_INT32": 32, "DT_UINT32": 32, "DT_INT64": 64, "DT_UINT64": 64,
    "DT_BOOL": 8, "DT_COMPLEX64": 64, "DT_COMPLEX128": 128,
})
# 仅兼容旧 --summary 模式；通用 --batch-dir 模式动态扫描，完全不使用此映射。
MODEL_FILES = {
    1: "slash_0/sg_aec_model_eef4b5dcb846_fuse_20261020_154551.omc",
    2: "slash_0/sg_Concat_1_58175fb15a27_fuse_20261020_154725.omc",
    3: "slash_0/sg_Concat_2_d99585f3b16f_fuse_20261020_154737.omc",
    4: "slash_0/sg_Concat_4d27b8b2be5c_fuse_20261020_154749.omc",
    5: "slash_0/sg_dual_model_f3db236e3a66_fixed_fuse_20261021_114854.omc",
    6: "slash_0/sg_Gemm_1903_1365bd2c80ad_fuse_20261020_154803.omc",
    7: "slash_0/sg_Gemm_3143_a04f843cfb6d_fuse_20261020_154815.omc",
    8: "slash_0/sg_Gemm_3606_cc457dee360e_fuse_20261020_154827.omc",
    9: "slash_0/sg_iccrn_model_f046edf3caed_fixed_fuse_20261021_115151.omc",
    10: "slash_0/sg_mono_model_da1e866a9015_fuse_20261020_154840.omc",
    11: "slash_0/sg_Split_1_eb7b7c75d820_fuse_20261020_154926.omc",
    12: "slash_0/sg_Split_32afaa784333_fuse_20261020_154938.omc",
    13: "slash_1/sg_aec_model_gate_conv1_636db64d27b2_fuse_20261020_154950.omc",
    14: "slash_1/sg_aec_model_gate_conv2_9168c58e49a1_fuse_20261020_155017.omc",
    15: "slash_1/sg_aec_model_gate_conv3_cdef48a44857_fuse_20261020_155049.omc",
    16: "slash_1/sg_aec_model_gate_conv4_67d1a778d921_fuse_20261020_155129.omc",
    17: "slash_1/sg_aec_model_gate_deconv3_1_5a94eef8de78_fuse_20261020_155211.omc",
    18: "slash_1/sg_aec_model_gate_deconv4_1_db9241d9053a_fuse_20261020_155228.omc",
    19: "slash_1/sg_aec_model_middle_ft_lstm1_a5f526137075_fuse_20261020_155245.omc",
    20: "slash_1/sg_dual_model_fb_8a856f6b6b56_fuse_20261020_155309.omc",
    21: "slash_1/sg_dual_model_freq_compress_e8184bc058d0_fuse_20261020_155337.omc",
    22: "slash_1/sg_dual_model_freq_decompress_3bc3d32555a9_fuse_20261020_155400.omc",
    23: "slash_1/sg_dual_model_oSpatialNet_freq_cmps_1_1a6f76e2b0c0_fixed_fuse_20261021_115152.omc",
    24: "slash_1/sg_iccrn_model_fb_223b058bb0c1_fixed_fuse_20261021_115154.omc",
    25: "slash_1/sg_iccrn_model_freq_compress_77652216f921_fuse_20261020_155412.omc",
    26: "slash_1/sg_iccrn_model_freq_decompress_292264a7d8e8_fuse_20261020_155429.omc",
    27: "slash_1/sg_iccrn_model_oSpatialNet_freq_cmps_1_3b2388043deb_fixed_fuse_20261021_115155.omc",
    28: "slash_1/sg_mono_model_cfb_d1_67793e29718f_fuse_20261020_155452.omc",
    29: "slash_1/sg_mono_model_cfb_d3_ab78bb99df64_fuse_20261020_155517.omc",
    30: "slash_1/sg_mono_model_in_ch_lstm1_a283428b8ac3_fuse_20261020_155540.omc",
    31: "slash_1/sg_mono_model_in_ch_lstm2_97af55dd52df_fuse_20261020_155558.omc",
    32: "slash_2/sg_aec_model_gate_conv1_u_block_65d1b9ecefd7_fuse_20261020_155616.omc",
    33: "slash_2/sg_aec_model_gate_conv2_u_block_6b3c2ed38a9f_fuse_20261020_155649.omc",
    34: "slash_2/sg_aec_model_gate_conv3_u_block_0bb706334f7e_fuse_20261020_155725.omc",
    35: "slash_2/sg_aec_model_gate_conv4_u_block_374f4b8ffec0_fuse_20261020_155801.omc",
    36: "slash_2/sg_aec_model_gate_deconv4_1_deconv2_de44dfb7649e_fuse_20261020_155838.omc",
    37: "slash_2/sg_aec_model_middle_ft_lstm1_Split_0733bada9736_fuse_20261020_155852.omc",
    38: "slash_2/sg_dual_model_fb_fullband_mamba_0_1_33655034b17f_fuse_20261020_155905.omc",
    39: "slash_2/sg_dual_model_freq_decompress_decompress_f9e7516e5b72_fuse_20261020_155926.omc",
    40: "slash_2/sg_dual_model_oSpatialNet_freq_cmps_1_mhsa_1_12f89ecc129e_fuse_20261020_155938.omc",
    41: "slash_2/sg_dual_model_oSpatialNet_freq_cmps_1_norm_full_1_9e85a83c6f82_fuse_20261020_155954.omc",
    42: "slash_2/sg_iccrn_model_fb_full_d836678d6958_fuse_20261020_160008.omc",
    43: "slash_2/sg_iccrn_model_fb_mhsa_1_fcb898dad124_fuse_20261020_160021.omc",
    44: "slash_2/sg_iccrn_model_fb_norm_full_af1e49f1c47d_fuse_20261020_160037.omc",
    45: "slash_2/sg_iccrn_model_oSpatialNet_freq_cmps_1_mhsa_1_5f990b089e7b_fuse_20261020_160050.omc",
    46: "slash_2/sg_iccrn_model_oSpatialNet_freq_cmps_1_norm_full_b30db64d67b0_fuse_20261020_160107.omc",
    47: "slash_2/sg_mono_model_cfb_d1_ceps_unit_bedf9c28ef19_fuse_20261020_160120.omc",
    48: "slash_2/sg_mono_model_in_ch_lstm1_lstm_cell_9a5934dbcec1_fuse_20261020_160139.omc",
    49: "slash_2/sg_mono_model_in_ch_lstm2_lstm_cell_a88fcd1007b3_fuse_20261020_160156.omc",
    50: "slash_3/sg_aec_model_gate_conv1_u_block_u2_deconv1_b818514c0173_fuse_20261020_160214.omc",
    51: "slash_3/sg_aec_model_gate_conv1_u_block_u2_deconv2_0d6ac9725f1c_fuse_20261020_160231.omc",
    52: "slash_3/sg_aec_model_gate_deconv4_1_deconv2_ConvTranspose_93753c142662_fuse_20261020_160248.omc",
    53: "slash_3/sg_dual_model_freq_decompress_decompress_MatMul_505218e9e784_fuse_20261020_160303.omc",
    54: "slash_3/sg_iccrn_model_fb_full_MatMul_d9e6019ea04b_fuse_20261020_160315.omc",
    55: "slash_3/sg_mono_model_cfb_d1_ceps_unit_irfft_b46ccf83f375_fuse_20261020_160326.omc",
    56: "slash_3/sg_mono_model_cfb_d1_ceps_unit_LN_19f0c43588bb_fuse_20261020_160339.omc",
}


def model_key(filename):
    """保留子图名字及哈希，只去掉编译/修复后缀，兼容两种路径分隔符。"""
    stem = PurePosixPath(filename.strip().replace("\\", "/")).stem
    stem = re.sub(r"_fuse_\d{8}_\d{6}$", "", stem)
    return re.sub(r"_fixed$", "", stem)


def read_power(path):
    if load_workbook is None:
        raise ValueError("缺少 openpyxl，请先运行：python -m pip install openpyxl")
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet_name = next((name for name in (SHEET_NAME, FALLBACK_SHEET_NAME)
                               if name in workbook.sheetnames), None)
            if sheet_name is None:
                raise ValueError(f"{path.name}: 找不到 sheet {SHEET_NAME!r} 或 {FALLBACK_SHEET_NAME!r}")
            sheet = workbook[sheet_name]
            # 某些导出器的 worksheet dimension 不准确，按指定区域实际读取。
            sheet.reset_dimensions()
            records = list(sheet.iter_rows(min_row=2, max_row=14, max_col=6))
            headers = [str(cell.value or "").strip() for cell in records[0]]
            expected = {"DB", "Inst", "Channel", "Current", "Voltage", "Power"}
            if len(set(headers)) != 6 or set(headers) != expected:
                raise ValueError(f"{path.name}: 第 2 行表头应为 {sorted(expected)}，实际为 {headers}")
            channel_col, power_col = headers.index("Channel"), headers.index("Power")
            powers = {}
            for row_number, cells in enumerate(records[1:], start=3):
                if all(cell.value is None for cell in cells):
                    continue
                channel = cells[channel_col].value
                if not isinstance(channel, str) or not channel.strip():
                    raise ValueError(f"{path.name}: 第 {row_number} 行 Channel 为空或不是文本")
                channel = channel.strip()
                if channel in powers:
                    raise ValueError(f"{path.name}: Channel {channel!r} 重复，不能确定唯一 Power")
                cell = cells[power_col]
                power = cell.value
                if power is None or (isinstance(power, str) and not power.strip()):
                    raise ValueError(
                        f"{path.name}: {channel} 的 Power 为空；若含公式，请先用 Excel 计算并保存")
                if (cell.data_type == "e" or isinstance(power, bool)
                        or not isinstance(power, (int, float, str))
                        or (isinstance(power, float) and not math.isfinite(power))):
                    raise ValueError(f"{path.name}: {channel} 的 Power 无效: {power!r}")
                powers[channel] = power
            if not powers:
                raise ValueError(f"{path.name}: 第 3～14 行没有 Channel/Power 数据")
            return powers
        finally:
            workbook.close()


def read_csv(path, required):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        fields = reader.fieldnames
        if not fields or any(not field.strip() for field in fields) or len(set(fields)) != len(fields):
            raise ValueError(f"{path}: CSV 表头为空，或存在重复/空列名")
        missing = set(required) - set(fields)
        if missing:
            raise ValueError(f"{path}: CSV 缺少列：{sorted(missing)}")
        rows = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"{path}: CSV 第 {reader.line_num} 行的字段数量与表头不一致")
            rows.append(row)
    return fields, rows


def same_value(left, right):
    if left == right:
        return True
    try:
        return Decimal(str(left)) == Decimal(str(right))
    except InvalidOperation:
        return False


def normalized_prefix(prefix):
    return "/".join(part for part in prefix.strip().split("/") if part)


def tensor_bytes(tensor):
    shape, dtype = tensor.get("shape"), tensor.get("dtype")
    if (not isinstance(shape, list) or any(type(size) is not int or size < 0 for size in shape)
            or type(dtype) not in (int, str) or dtype not in DTYPE_BITS):
        return None
    return (math.prod(shape) * DTYPE_BITS[dtype] + 7) // 8


def transfer_statistics(report):
    """按 stats.json 的节点成员归并，每个层级独立；整图只累计各节点一次。"""
    nodes = report.get("nodes")
    if not isinstance(nodes, list) or len(nodes) != report["total"]["node_count"]:
        raise ValueError("stats.json 缺少完整 nodes 元数据，请用 stats.py 生成完整报告")
    per_node = {}
    for node in nodes:
        node_id = node["node_id"]
        if type(node_id) is not int or node_id in per_node:
            raise ValueError(f"stats.json 存在重复或无效节点 ID：{node_id!r}")
        values, known = {}, 0
        for direction, field in (("inputs", "输入及权重字节数"), ("outputs", "输出字节数")):
            tensors = node.get(direction)
            if not isinstance(tensors, list) or any(not isinstance(t, dict) for t in tensors):
                raise ValueError(f"stats.json 节点 {node_id} 缺少 {direction} 张量元数据")
            sizes = [tensor_bytes(t) for t in tensors
                     if not (report.get("source_format") == "OMC" and direction == "inputs"
                             and str(t.get("name", "")).endswith(":-1"))]
            values[field] = sum(sizes) if all(size is not None for size in sizes) else None
            known += sum(size for size in sizes if size is not None)
        values["数据转移字节数"] = (sum(values.values()) if all(v is not None for v in values.values()) else None)
        values["已知数据转移字节数"] = known
        per_node[node_id] = values

    def totals(ids):
        if len(set(ids)) != len(ids) or any(node_id not in per_node for node_id in ids):
            raise ValueError("stats.json 的 node_ids 重复或引用了不存在的节点")
        members = [per_node[node_id] for node_id in ids]
        result = {}
        for field in TRANSFER_FIELDS[:3]:
            values = [node[field] for node in members]
            result[field] = sum(values) if all(value is not None for value in values) else None
        result["未知数据转移节点数"] = sum(node["数据转移字节数"] is None for node in members)
        result["数据转移量完整"] = "否" if result["未知数据转移节点数"] else "是"
        result["已知数据转移字节数"] = sum(node["已知数据转移字节数"] for node in members)
        return result

    groups = {}
    for level in report["levels"]:
        for group in level["subgraphs"]:
            key = (level["slash_count"], normalized_prefix(group["prefix"]))
            if key in groups or len(group["node_ids"]) != group["node_count"]:
                raise ValueError(f"stats.json 子图重复或节点数量不一致：{key}")
            groups[key] = {"节点数": group["node_count"], **totals(group["node_ids"])}
    return groups, totals(list(per_node))


def full_model_statistics(stats_dir, fields, report=None):
    path = stats_dir / "stats.json"
    try:
        if report is None:
            report = json.loads(path.read_text(encoding="utf-8-sig"))
        total = report["total"]
        result = {
            "ONNX文件": report["source"], "节点数": total["node_count"],
            "MACs": total["macs"], "FLOPs": total["flops"],
            "MFLOPs": total["flops"] / 1e6 if total["flops"] is not None else None,
            "计算量完整": "是" if total["complete"] else "否",
            "未知计算量节点数": len(total["unknown_node_ids"]),
            "已知MACs": total["known_macs"], "已知FLOPs": total["known_flops"],
        }
        result.update({field: total["operator_counts"].get(field[:-2], 0)
                       for field in fields if field.endswith("数量")})
        if "exec_time" in total:
            result.update({
                "exec_time": total["exec_time"], "exec_time_unit": "us",
                "算术强度": total["arithmetic_intensity"], "算术强度单位": "FLOPs/us",
                "已知exec_time": total["known_exec_time"],
                "时间匹配节点数": total["exec_time_matched_nodes"],
                "时间未匹配节点数": len(total["exec_time_missing_node_ids"]),
            })
    except (KeyError, TypeError, AttributeError) as error:
        raise ValueError(f"{path}: 不符合 stats.py 生成的统计报告格式（{error}）") from error
    return {field: value for field, value in result.items() if field in fields}


def merge_statistics(fields, rows, stats_dir):
    stats_dir = Path(stats_dir)
    if "斜杠数" not in fields:
        raise ValueError("合并统计时，原总表必须包含 斜杠数 列")
    paths = sorted((int(match[1]), path) for path in stats_dir.glob("stats_*_slash.csv")
                   if path.is_file() and (match := re.fullmatch(r"stats_(\d+)_slash\.csv", path.name)))
    if not paths:
        raise FileNotFoundError(f"{stats_dir}: 找不到 stats_<数字>_slash.csv")
    statistics, stat_fields = {}, []
    for depth, path in paths:
        columns, records = read_csv(path, ("子图名字", "节点数", "FLOPs"))
        stat_fields.extend(column for column in columns
                           if column != "子图名字" and column not in stat_fields)
        for record in records:
            prefix = normalized_prefix(record["子图名字"])
            if not prefix or prefix.count("/") != depth:
                raise ValueError(f"{path.name}: 子图名字 {record['子图名字']!r} 与斜杠层级不一致")
            key = (depth, prefix)
            if key in statistics:
                raise ValueError(f"统计表重复出现同一子图：slash_{depth}/{prefix}")
            statistics[key] = record

    json_path = stats_dir / "stats.json"
    if not json_path.is_file():
        raise FileNotFoundError(f"{json_path}: 数据转移量需要 stats.py 生成的完整 stats.json（含 nodes、levels）")
    try:
        report = json.loads(json_path.read_text(encoding="utf-8-sig"))
        transfers, full_transfer = transfer_statistics(report)
    except (KeyError, TypeError, AttributeError) as error:
        raise ValueError(f"{json_path}: 缺少数据转移量所需的节点/子图元数据（{error}）") from error
    for row in rows:
        prefix = normalized_prefix(row["子图名字"])
        key = (int(row["斜杠数"]), prefix)
        if key not in statistics or key not in transfers:
            raise ValueError(f"统计 CSV 或 stats.json 找不到总表子图：slash_{key[0]}/{prefix}")
        source = dict(statistics[key])
        for field, value in transfers[key].items():
            if field in source and not same_value(source[field], "" if value is None else value):
                raise ValueError(f"子图 {prefix!r} 的 {field} 在统计 CSV 与 stats.json 中不一致")
            source[field] = value
        for field in fields:
            if field != "子图名字" and field in source and not same_value(row[field], source[field]):
                raise ValueError(
                    f"子图 {prefix!r} 的 {field} 冲突：总表 {row[field]!r}，"
                    f"统计表 {source[field]!r}；请使用同一批统计数据")
        row.update({field: source.get(field, "") for field in [*stat_fields, *TRANSFER_FIELDS]
                    if field not in fields})

    output_fields = list(dict.fromkeys([*fields, *stat_fields, *TRANSFER_FIELDS]))
    if "算术强度" not in stat_fields:
        print("提示：源统计表没有算术强度列，请用带 profiling 的统计表；此次仅合并已有统计。", file=sys.stderr)
    full_stats = full_model_statistics(stats_dir, output_fields, report)
    full_stats.update(full_transfer)
    return output_fields, full_stats


def merge_power(summary_path, power_dir, output_path=None, *, stats_dir=None):
    summary_path, power_dir = Path(summary_path), Path(power_dir)
    output_path = (Path(output_path) if output_path else
                   summary_path.with_name(summary_path.stem + (
                       "_with_power_stats.csv" if stats_dir is not None else "_with_power.csv")))
    if output_path.exists():
        raise FileExistsError(f"输出文件已存在，请用 -o 指定新文件：{output_path}")

    fields, rows = read_csv(summary_path, ("ONNX文件", "子图名字"))
    model_ids = {model_key(name): number for number, name in MODEL_FILES.items()}
    if set(MODEL_FILES) != set(range(1, 57)) or len(model_ids) != 56:
        raise ValueError("内置编号映射不完整或存在重复子图，请检查 MODEL_FILES")
    matched_rows = {}
    for line_number, row in enumerate(rows, start=2):
        number = model_ids.get(model_key(row["ONNX文件"]))
        if number is None:
            raise ValueError(f"CSV 第 {line_number} 行无法匹配编号：{row['ONNX文件']!r}")
        if number in matched_rows:
            raise ValueError(f"CSV 中同一子图重复匹配到 {number}.xlsx：{row['ONNX文件']!r}")
        matched_rows[number] = row
    missing = sorted(set(MODEL_FILES) - matched_rows.keys())
    if missing:
        raise ValueError(f"CSV 缺少以下功耗编号对应的子图：{missing}")

    full_stats = {}
    if stats_dir is not None:
        fields, full_stats = merge_statistics(fields, rows, stats_dir)

    paths = {number: power_dir / f"{number}.xlsx" for number in range(57)}
    missing_files = [path.name for path in paths.values() if not path.is_file()]
    if missing_files:
        raise FileNotFoundError(f"功耗目录缺少文件：{', '.join(missing_files)}")
    measurements = {number: read_power(path) for number, path in paths.items()}
    channels = list(measurements[0])
    collisions = set(channels) & set(fields)
    if collisions:
        raise ValueError(f"CSV 已有同名 Channel 列，拒绝覆盖：{sorted(collisions)}")
    for number, powers in measurements.items():
        if set(powers) != set(channels):
            raise ValueError(
                f"{number}.xlsx 的 Channel 与 0.xlsx 不一致；"
                f"缺少 {sorted(set(channels) - powers.keys())}，"
                f"多出 {sorted(powers.keys() - set(channels))}")

    for number, row in matched_rows.items():
        row.update(measurements[number])
    full_model = dict.fromkeys(fields, "")
    full_model.update(full_stats)
    full_model["子图名字"] = "整模型"
    full_model.update(measurements[0])
    rows.append(full_model)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields + channels)
        writer.writeheader()
        writer.writerows(rows)
    return output_path, channels


def _power_csv_rows(path):
    raw = Path(path).read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("gb18030")
    import io
    return list(csv.reader(io.StringIO(text)))


def read_runtime_seconds(path):
    """读取 CSV 第一行第八列或选中平均值 sheet 的 H1；缺失返回 None，单位为秒。"""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        rows = _power_csv_rows(path)
        value = rows[0][7] if rows and len(rows[0]) > 7 else None
    else:
        if load_workbook is None:
            raise ValueError("读取 XLSX 时长需要 openpyxl")
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            try:
                name = next((n for n in (SHEET_NAME, FALLBACK_SHEET_NAME) if n in book.sheetnames), None)
                if name is None:
                    raise ValueError(f"{path.name}: 找不到平均值 sheet，无法读取 H1")
                sheet = book[name]
                sheet.reset_dimensions()
                cell = sheet.cell(row=1, column=8)
                if cell.data_type == "e":
                    raise ValueError(f"{path.name}: H1 为 Excel 错误值")
                value = cell.value
            finally:
                book.close()
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError(f"{path.name}: H1 时长不是有效数值")
    try:
        seconds = Decimal(str(value).strip())
    except InvalidOperation as error:
        raise ValueError(f"{path.name}: H1 时长不是秒数：{value!r}") from error
    if not seconds.is_finite() or seconds < 0:
        raise ValueError(f"{path.name}: H1 时长必须为有限非负秒数")
    result = float(seconds)
    if not math.isfinite(result) or (seconds != 0 and result == 0):
        raise ValueError(f"{path.name}: H1 时长超出可表示范围")
    return result


def read_power_csv(path):
    """沿用平均值表的 Channel/Power 及前 12 行；保留单位，不对重复 Channel 猜测求均值。"""
    rows = _power_csv_rows(path)
    headers = []
    for index, row in enumerate(rows):
        fields = [str(x).strip().lower() for x in row]
        while fields and not fields[-1]:
            fields.pop()
        if {"channel", "power"}.issubset(fields):
            if len(set(fields)) != len(fields) or "" in fields:
                raise ValueError(f"{path.name}: CSV 表头有重复或空字段")
            headers.append((index, fields))
    if not headers:
        raise ValueError(f"{path.name}: 找不到包含 Channel、Power 的 CSV 表头")
    blocks = []
    for n, (index, fields) in enumerate(headers):
        end = headers[n + 1][0] if n + 1 < len(headers) else len(rows)
        records = [r for r in rows[index + 1:end] if any(x.strip() for x in r)][:12]
        if not records:
            continue
        db = fields.index("db") if "db" in fields else None
        marker = " ".join(r[db] for r in records if db is not None and len(r) > db).lower()
        if not marker and index:
            marker = " ".join(rows[index - 1]).lower()
        rank = 0 if "work0.db" in marker else 1 if "work1.db" in marker else 2
        blocks.append((rank, fields, records))
    if not blocks:
        raise ValueError(f"{path.name}: CSV 没有 Channel/Power 数据")
    rank = min(b[0] for b in blocks)
    selected = [b for b in blocks if b[0] == rank]
    if len(selected) != 1:
        raise ValueError(f"{path.name}: 存在多个同优先级功耗表，不能确定唯一数据")
    _, fields, records = selected[0]
    c, p = fields.index("channel"), fields.index("power")
    powers = {}
    for record in records:
        if len(record) <= max(c, p):
            raise ValueError(f"{path.name}: CSV 数据行缺少 Channel 或 Power")
        channel, power = record[c].strip(), record[p].strip()
        if not channel or channel in powers:
            raise ValueError(f"{path.name}: Channel 为空或重复：{channel!r}")
        if not power or power.lower() in {"nan", "inf", "+inf", "-inf", "infinity", "#div/0!", "#value!"}:
            raise ValueError(f"{path.name}: {channel} 的 Power 为空或无效：{power!r}")
        powers[channel] = power
    return powers


def merge_compiled_power(batch_dir, power_dir, output_dir, *, ge_proto_path=None, numbering_csv=None):
    """动态编号、OMC 整图统计、同号功耗合并；缺失功耗保留空值及原因。"""
    if __package__:
        from .compiled_stats import analyze_compiled_subgraphs, number_models
    else:
        from compiled_stats import analyze_compiled_subgraphs, number_models
    from collections import Counter
    root, output = Path(batch_dir).resolve(), Path(output_dir).resolve()
    powers_root = Path(power_dir).resolve() if power_dir is not None else None
    if powers_root is not None and not powers_root.is_dir():
        raise NotADirectoryError(powers_root)
    if output.exists():
        raise FileExistsError(f"输出目录已存在，请指定新目录：{output}")
    numbering = number_models(root, numbering_csv)
    compiled = analyze_compiled_subgraphs(root, output / "omc_stats", ge_proto_path=ge_proto_path)
    by_path = {Path(r["OMC文件"]).resolve(): r for r in compiled["records"]}
    with (output / "numbering.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["编号", "斜杠数", "OMC相对路径", "OMC_SHA256"])
        writer.writeheader()
        writer.writerows(numbering)
    rows, channels, unused_power = [], [], []
    for item in numbering:
        row = dict(by_path[(root / item["OMC相对路径"]).resolve()])
        if row["OMC状态"] == "ok" and row["OMC_SHA256"] != item["OMC_SHA256"]:
            raise ValueError(f"统计期间 OMC 文件发生变化：{item['OMC相对路径']}")
        row.update(item)
        row.update(功耗文件="", 功耗状态="missing", 功耗原因="尚未提供功耗目录")
        row.update(运行时长=None, 运行时长单位="s", 时长状态="missing", 时长原因="尚未提供功耗目录",
                   数据移动吞吐=None, 数据移动吞吐单位="B/s")
        if powers_root is not None:
            candidates = [p for suffix in (".csv", ".xlsx") if (p := powers_root / f"{item['编号']}{suffix}").is_file()]
            if not candidates:
                row["功耗原因"] = f"缺少 {item['编号']}.csv（或 .xlsx），未知功耗不补零"
                row["时长原因"] = f"缺少 {item['编号']}.csv（或 .xlsx），未读取 H1"
            elif len(candidates) > 1:
                row.update(功耗状态="ambiguous", 功耗原因="同一编号同时存在 CSV 和 XLSX，请保留一个对应文件")
                row.update(时长状态="ambiguous", 时长原因="同一编号同时存在 CSV 和 XLSX，无法确定时长来源")
            else:
                path = candidates[0]
                row["功耗文件"] = str(path)
                try:
                    values = read_power_csv(path) if path.suffix.lower() == ".csv" else read_power(path)
                    if set(values) & set(row):
                        raise ValueError(f"Channel 与统计列同名：{sorted(set(values) & set(row))}")
                    channels.extend(c for c in values if c not in channels)
                    row.update(values)
                    row.update(功耗状态="ok", 功耗原因="")
                except (OSError, ValueError, BadZipFile, InvalidFileException) as error:
                    row.update(功耗状态="failed", 功耗原因=str(error))
                try:
                    seconds = read_runtime_seconds(path)
                    if seconds is None:
                        row["时长原因"] = "H1 为空或缺少第一行第八列；若为公式，请先计算并保存"
                    else:
                        row.update(运行时长=seconds, 时长状态="ok", 时长原因="")
                        amount = row.get("OMC_数据转移字节数")
                        if amount is not None and amount != "" and seconds > 0:
                            rate = amount / seconds
                            row["数据移动吞吐"] = rate if math.isfinite(rate) else None
                except (OSError, ValueError, BadZipFile, InvalidFileException) as error:
                    row.update(时长状态="failed", 时长原因=str(error))
        rows.append(row)
    for row in rows:
        if row["功耗状态"] == "ok" and any(c not in row for c in channels):
            row.update(功耗状态="incomplete", 功耗原因="缺少 Channel：" + ", ".join(c for c in channels if c not in row))
    if powers_root is not None:
        expected = {f"{r['编号']}{ext}" for r in numbering for ext in (".csv", ".xlsx")}
        unused_power = [str(p) for p in sorted(powers_root.iterdir())
                        if p.is_file() and p.suffix.lower() in {".csv", ".xlsx"} and p.name.lower() not in expected]
    base_fields = list(dict.fromkeys(k for row in rows for k in row if k not in channels and k != "operator_counts"))
    headers = ["编号", "斜杠数", "OMC相对路径", *[f for f in base_fields if f not in {"编号", "斜杠数", "OMC相对路径"}], *channels]
    with (output / "summary.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    result = dict(batch_dir=str(root), power_dir=str(powers_root) if powers_root is not None else None,
                  source_format="OMC", numbering_rule="slash 数字升序、文件名忽略大小写升序；从 1 开始",
                  omc_counts=compiled["counts"], power_counts=dict(Counter(r["功耗状态"] for r in rows)),
                  runtime_counts=dict(Counter(r["时长状态"] for r in rows)),
                  unused_power_files=unused_power, records=rows)
    (output / "summary.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--batch-dir", type=Path, help="通用模式：包含 slash_0/slash0 等子目录的 OMC 文件夹")
    mode.add_argument("--summary", type=Path, help="旧模式：代表子图 representatives_summary.csv")
    parser.add_argument("--power-dir", type=Path, help="通用模式读同号 N.csv/N.xlsx；旧模式读 0.xlsx～56.xlsx")
    parser.add_argument("--output-dir", type=Path, help="通用模式输出目录，必须尚不存在")
    parser.add_argument("--ge-proto", type=Path, help="通用模式：DDK 的 ge_ir_pb2.py")
    parser.add_argument("--numbering", type=Path, help="通用模式：复用已生成的 numbering.csv，校验文件 SHA256")
    parser.add_argument("--stats-dir", type=Path, help="统计目录：stats_<数字>_slash.csv 和完整 stats.json")
    parser.add_argument("-o", "--output", type=Path, help="输出 CSV；默认后缀为 _with_power 或 _with_power_stats")
    args = parser.parse_args()
    if args.batch_dir is not None:
        if args.output is not None or args.stats_dir is not None:
            parser.error("--batch-dir 模式请用 --output-dir；OMC 统计自动生成，不使用 --stats-dir")
        output = args.output_dir or args.batch_dir / "omc_power_summary"
        try:
            result = merge_compiled_power(args.batch_dir, args.power_dir, output,
                                          ge_proto_path=args.ge_proto, numbering_csv=args.numbering)
        except (OSError, ValueError) as error:
            print(f"合并失败：{error}", file=sys.stderr)
            return 1
        print(f"OMC: {result['omc_counts']}；功耗: {result['power_counts']}；时长: {result['runtime_counts']}")
        print(f"编号表：{Path(output).resolve() / 'numbering.csv'}")
        print(f"总表：{Path(output).resolve() / 'summary.csv'}")
        return 0
    if args.power_dir is None:
        parser.error("--summary 旧模式需要 --power-dir；通用 OMC 模式使用 --batch-dir")
    try:
        output, channels = merge_power(args.summary, args.power_dir, args.output, stats_dir=args.stats_dir)
    except (OSError, ValueError, BadZipFile, InvalidFileException) as error:
        print(f"合并失败：{error}", file=sys.stderr)
        return 1
    print(f"合并完成：56 个子图 + 1 个整模型，新增 {len(channels)} 个 Channel 功耗列。")
    print(f"Channel：{', '.join(channels)}")
    if args.stats_dir is not None:
        print(f"已合并统计和数据转移量：{args.stats_dir.resolve()}")
    print(f"输出：{output.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
